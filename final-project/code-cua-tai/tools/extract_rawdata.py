"""Extract text from raw course files into reviewable JSON/JSONL artifacts."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
from pathlib import Path

from docx import Document
from pypdf import PdfReader

SUPPORTED = {".pdf", ".docx", ".doc", ".txt"}
LESSON_DOCUMENTS = {
    "B1": "1. Tong quan ve TTNT.pdf",
    "B2": "2. Thuat giai Heuristic.pdf",
    "B3": "3. Phuong phap tim kiem.pdf",
    "B4": "4. Chien luoc Minimax.doc",
    "B5": "5. Tong quan ve BDTT.pdf",
    "B6a": "6a. Cac phuong phap BDTT co ban.pdf",
}

def clean_text(value: str) -> str:
    value = value.encode("utf-8", errors="replace").decode("utf-8")
    value = value.replace("\x00", " ").replace("﻿", " ")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()

def extract_pdf(path: Path) -> tuple[str, int]:
    reader = PdfReader(str(path))
    pages = [clean_text(page.extract_text() or "") for page in reader.pages]
    return "\n\n".join(text for text in pages if text), len(reader.pages)

def extract_docx(path: Path) -> tuple[str, int]:
    document = Document(str(path))
    paragraphs = [clean_text(paragraph.text) for paragraph in document.paragraphs]
    tables = [" | ".join(clean_text(cell.text) for cell in row.cells) for table in document.tables for row in table.rows]
    return "\n".join(text for text in paragraphs + tables if text), len(document.paragraphs)

def extract_text(path: Path) -> tuple[str, int]:
    return clean_text(path.read_text(encoding="utf-8", errors="replace")), 1

def extract_doc(path: Path) -> tuple[str, int]:
    result = subprocess.run(["antiword", str(path)], capture_output=True, check=True)
    return clean_text(result.stdout.decode("utf-8", errors="replace")), 1

def extract_file(path: Path) -> tuple[str, int]:
    if path.suffix.lower() == ".pdf":
        return extract_pdf(path)
    if path.suffix.lower() == ".docx":
        return extract_docx(path)
    if path.suffix.lower() == ".doc":
        return extract_doc(path)
    if path.suffix.lower() == ".txt":
        return extract_text(path)
    return "", 0

def main() -> None:
    parser = argparse.ArgumentParser(description="Extract rawdata documents for knowledge-base review")
    parser.add_argument("--input", type=Path, default=Path("rawdata"))
    parser.add_argument("--output", type=Path, default=Path("data/rawdata-extracted.jsonl"))
    args = parser.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    records = []
    errors = []
    for path in sorted(args.input.iterdir()):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED:
            continue
        try:
            text, units = extract_file(path)
            lesson = next((lesson_id for lesson_id, filename in LESSON_DOCUMENTS.items() if filename == path.name), None)
            records.append({
                "id": hashlib.sha256(path.name.encode("utf-8")).hexdigest()[:16],
                "file": path.name,
                "format": path.suffix.lower().lstrip("."),
                "lesson": lesson,
                "units": units,
                "characters": len(text),
                "text": text,
            })
        except Exception as exc:  # keep the batch running and report the bad file
            errors.append({"file": path.name, "error": str(exc)})
    with args.output.open("w", encoding="utf-8", newline="\n") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False) + "\n")
    summary = {
        "files": len(records),
        "filesWithText": sum(bool(record["text"]) for record in records),
        "characters": sum(record["characters"] for record in records),
        "errors": errors,
        "output": str(args.output),
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))

if __name__ == "__main__":
    main()
